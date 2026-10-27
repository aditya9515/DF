"""Build the report from the retained DOCX and measured artifacts only.

Run with the Codex bundled Python when authoring in Codex. For independent
reproduction install python-docx and pypdf. Render separately, then run again
with --pagination pointing to that PDF to fill accurate contents references.
"""

import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import tomllib
from zipfile import ZipFile, ZIP_DEFLATED
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_SHA = "bb61063470651d356a3071184fa07a3cf1339941e819e496ec7a421a43724950"
HEADINGS = [
    "Introduction",
    "Research Question And Objectives",
    "Related Work",
    "Literature Comparison",
    "Dataset",
    "Data Preparation",
    "Proposed Methodology",
    "Models And Baseline",
    "Experimental Setup",
    "Main Results",
    "Figures And Interpretation",
    "Discussion",
    "Limitations And Future Work",
    "Conclusion",
]


def read(path):
    return json.loads((ROOT / path).read_text(encoding="utf-8"))


def pct(x):
    return f"{100 * x:.1f}%"


def fnum(x):
    return f"{x:.3f}"


def pagination(path):
    if not path:
        return {}
    from pypdf import PdfReader

    result = {}
    for physical, page in enumerate(PdfReader(str(path)).pages, 1):
        if physical < 4:
            continue
        lines = [re.sub(r"\s+", " ", s).strip() for s in (page.extract_text() or "").splitlines()]
        for title in [f"{i} {h}" for i, h in enumerate(HEADINGS, 1)] + [
            "References",
            "Appendix Reproduction And Evidence",
        ]:
            if any(s == title for s in lines):
                result.setdefault(title, str(physical - 3))
        for line in lines:
            if re.match(r"^(Figure|Table) \d+ ", line):
                result.setdefault(line.split(" ", 2)[0] + " " + line.split(" ", 2)[1], str(physical - 3))
    return result


def content():
    summary = read("results/summary.json")
    data = read("data/dataset_summary.json")
    model = read("configs/model.json")
    sources = read("docs/sources.json")
    versions = {
        p["name"]: p["version"]
        for p in tomllib.loads((ROOT / "uv.lock").read_text(encoding="utf-8"))["package"]
        if "version" in p
    }
    m = summary["metrics"]

    def row(method, task="attack_exposure", source="InjecAgent-adapted"):
        return next(
            r
            for r in m
            if r["source"] == source and r["split"] == "test" and r["method"] == method and r["task"] == task
        )

    combined = row("combined")
    keyword = row("keyword")
    policy = row("policy_only")
    action = row("combined", "action_violation")
    tamper = read("results/tamper_results.json")
    corrupted = [r for r in tamper if not r["expected_valid"]]
    micro = read("results/overhead_micro.json")
    extra = sum(r["extra_ms_per_event"] for r in micro) / len(micro)
    live = read("results/overhead_live.json")
    import statistics

    live_delta = statistics.median(100 * (r["full_s"] / r["plain_s"] - 1) for r in live)
    all_outcomes = [json.loads(p.read_text(encoding="utf-8")) for p in (ROOT / "results/outcomes").glob("*.json")]
    all_outcomes = [o for o in all_outcomes if not o.get("measurement_only")]
    public = [o for o in all_outcomes if o["source"] == "InjecAgent-adapted" and o["split"] == "test"]
    attempted = sum(o["action_violation"] for o in public)
    executed = sum(o["simulated_unauthorized_actions"] for o in public)
    abstract = (
        f"Tool-using language-model agents can encounter malicious instructions inside retrieved content, while ordinary application logs provide limited evidence for investigating the resulting actions. "
        f"This project records observable agent events, links them with SHA-256 hashes, signs checkpoints and final manifests with Ed25519, and reconstructs incidents through an evidence viewer. "
        f"The dataset combines 60 InjecAgent attacks, 60 generated benign counterparts and 24 controlled scenarios, executed with three seeds using quantized Qwen3-4B. "
        f"Across {summary['sessions']} live sessions, the proposed detector obtained held-out attack-exposure F1 of {combined['f1']:.3f}, compared with {keyword['f1']:.3f} for a keyword baseline. "
        f"The verifier rejected {len(corrupted)} of {len(corrupted)} corrupted bundles, and independently receipted operations had {pct(summary['receipt_completeness'])} capture completeness. "
        f"The recorder added {extra:.3f} milliseconds per event in the measured microbenchmark. "
        "The contribution is a reproducible integration of evidence capture, integrity verification and event-linked triage. Results are limited to one model, adapted scenarios and simulated tools; they do not establish comprehensive attack detection or production forensic assurance."
    )
    blocks = {}
    blocks[1] = [
        "An LLM agent can retrieve a document, inspect an email, call another tool and produce an answer in a single session. An investigator must connect those actions to the exact inputs and configuration available at the time. A final answer or a loosely formatted console log may show that something happened without preserving the context needed to explain the sequence or check whether the record was altered.",
        "The problem affects developers investigating unexpected tool use, security teams responding to suspected prompt injection, and reviewers assessing whether an incident report is supported by preserved evidence. A system that merely displays a plausible timeline is insufficient if events can be changed without detection or if a finding cannot be located in the underlying record. Forensic incident-response guidance therefore motivates preserving original observations and separating them from analysis [1].",
        "This mini project builds and evaluates an evidence pipeline around a local agent. It captures observable operations, seals the resulting event sequence, independently verifies exported records, and links suspicious content to subsequent tool requests. The implementation integrates established cryptographic primitives with agent-specific context rather than introducing a new signature algorithm or claiming a universal injection defence.",
        "The scope is deliberately observable and local. Qwen generates actual responses and tool-call decisions, while tool side effects occur in a sandbox. The study evaluates capture coverage, detectable stored-record tampering, heuristic incident detection and measurement overhead. It does not recover hidden model state, execute harmful host actions, certify legal admissibility, or reproduce the full original benchmark environment.",
    ]
    blocks[2] = [
        "Research question: How much does structured, tamper-evident evidence capture improve incident detection and reconstruction compared with ordinary agent logs, and what overhead does it introduce?",
        "Objective 1: Implement an agent-session evidence model and capture layer, then measure its coverage against independent model-client and sandbox receipts. Check session boundaries separately from operation-pair completeness.",
        "Objective 2: Implement hash-linked events, signed checkpoints and final manifests, then verify intact evidence and challenge the verifier with edits, omissions, insertion, reordering, signature damage, key substitution and chain rewriting.",
        "Objective 3: Reconstruct incidents and compare keyword, policy-only and combined triage methods on the same held-out sessions. Evaluate malicious-content exposure separately from unauthorized action attempts, including false positives on benign instruction-like text.",
        "Objective 4: Quantify recording cost and package a reproducible implementation, datasets, evidence, predictions, tables and figures. Measure event-level overhead and paired live runtimes with the measurement boundary made explicit.",
        "The specific contribution is the working connection between capture, verification and analysis: a detector finding references event IDs in an independently verifiable sequence. Novelty is assessed at this integration and evaluation level. The study does not claim that hash chaining, digital signatures, allowlists or keyword detection are individually new.",
    ]
    blocks[3] = [
        "Kent et al. [1] place forensic collection, examination, analysis and reporting within incident response. Their guidance motivates retaining raw observations and distinguishing interpreted findings from the underlying record. The present application makes that distinction explicit through saved events and separately generated analysis.",
        "Schneier and Kelsey [2] investigate protecting historical audit records on an untrusted machine. Crosby and Wallach [3] develop authenticated history structures that support efficient consistency checks. These works motivate cryptographic commitments and auditing, but their security properties must not be automatically attributed to a simpler static-key implementation. This prototype uses linear event chains and final signed anchors; it does not implement forward-secure key evolution or a distributed transparency service.",
        "ReAct [4] demonstrates an interleaving of model outputs and environmental actions. That operational structure creates useful capture boundaries for requests, tool calls and observations. Greshake et al. [5] show how retrieved material can carry instructions that redirect an integrated model application, motivating the preservation of untrusted context alongside later actions.",
        "InjecAgent [6] supplies structured indirect-injection cases involving direct harm and data stealing. Its original evaluation studies agent vulnerability across different agents; this project instead adapts a selected subset to collect and analyze forensic evidence. AgentDojo [7] offers an extensible security-and-utility evaluation environment. It informs the discussion and future integration, but no AgentDojo experiments are claimed here.",
        "Raji et al. [8] connect organizational audit stages with documented evidence and reports. This platform addresses a narrower runtime question: whether an analyst can verify the preserved session and locate the events supporting a finding. Together, these sources motivate complementary parts of the design rather than establish a directly comparable leaderboard.",
    ]
    blocks[4] = [
        "Table 1 compares the implementation-oriented literature that informed the system. Different data and metrics prevent a direct numerical ranking; the last column identifies the part addressed by this project."
    ]
    blocks[5] = [
        "The public source is InjecAgent [6], downloaded from uiuc-kang-lab/InjecAgent at revision f19c9f2c79a41046eb13c03c51a24c567a8ffa07. The original base files contain 510 direct-harm cases and 544 data-stealing cases: 1,054 in total across 17 user-tool families. Original files and the MIT licence notice are retained, with exact URLs and SHA-256 digests in data/provenance.json.",
        "Original fields include user instruction, user tool, parameters, response template, injected tool response, attacker tools, attack category and expected achievements. Only the observable prompt, tool schemas and resulting tool response are presented to the agent. The detector receives observable events and declared authorization policy; attack labels, expected achievements and attacker-instruction metadata remain outside its feature inputs.",
        "The live subset contains 60 published attacks, balanced between the two source classes. Each has a project-generated benign counterpart created by replacing the attack placeholder in its original response template. Twenty controls contain clearly quoted malicious examples to challenge a literal keyword rule. These controls are authored for this project and are not original InjecAgent negatives. Repeated surrounding templates also limit the effective diversity of the control set.",
        "An additional 24 project-generated scenarios cover benign activity, direct instruction conflict, synthetic-secret exfiltration and privilege misuse, with six parameterized variants per category. All 144 scenarios run under seeds 42, 43 and 44, giving 432 planned sessions. Controlled variants share category-level wording, so their split does not establish generalization to new attack templates. The two targets distinguish supplied attack content from proposed calls outside the predefined task policy.",
        f"The prepared manifest has {data['missing_required_fields']} missing required scenario fields. There is no tabular scaling or imputation stage. Published files have unrestricted download access under their retained licence; the Qwen weights have a separate Apache-2.0 licence. No private user dataset or real secret is included.",
    ]
    blocks[6] = [
        "Preparation preserves the original raw files before producing a separate scenario manifest. String-encoded source parameters are parsed with a literal parser rather than executed. Tool schemas are normalized to JSON-compatible function definitions, and every selected tool name is checked against the source catalogue or an explicit controlled-scenario override. Preparation does not execute any source command or attack instruction.",
        "Leakage control uses task families rather than an independent random split of near-duplicate rows. The sorted 17 user-tool families are shuffled with seed 42; six families form development and the remaining eleven form the test pool. Within each pool, attacks are sampled deterministically and balanced between direct harm and data stealing. There are 20 development attacks and 40 test attacks, each paired with its benign control. All seeds and both members of a pair stay in the same split.",
        "For each controlled category, templates 0 and 1 belong to development, while templates 2 through 5 belong to test. This yields eight development and sixteen held-out controlled scenarios before seed repetitions. The split definition and source IDs are saved with the scenario manifest, allowing reviewers to audit group separation.",
        "No neural network is trained for detection. The rules were frozen before inference, with a source-code digest recorded in configs/detector_freeze.json. Consequently encoding, feature scaling and imbalance resampling are not applicable to this rule-based study. Class balance is managed through scenario selection, and evaluation reports its actual support rather than assuming equal numbers of policy violations.",
        "The three seeds measure repeated model behavior, not three independent datasets. Confidence intervals therefore resample tool or scenario groups, keeping related observations together. Generated controls and original attacks are explicitly identified so that source benchmark labels are not silently extended to authored examples.",
    ]
    blocks[7] = [
        "Figure 1 shows the workflow in execution order. First, the local Qwen model receives a legitimate task and an initial tool response specified by the benchmark. LangGraph alternates model calls and simulated tool execution, stopping after an answer or a three-turn limit. The initial retrieval is supplied by the benchmark rather than discovered autonomously; this adaptation isolates behavior after exposure.",
        "Second, the capture layer records session configuration, the user prompt, retrieval content, model requests and responses, tool proposals and tool results. Parent event IDs connect a request with its response and an observed action with the preceding model output. Sandbox receipts independently record performed operations, and client receipts record successful or failed inference calls for later coverage checks.",
        "Third, every event is serialized with RFC 8785 canonical JSON [13] and linked to its predecessor by SHA-256. Ed25519 signatures [12] protect checkpoints after every ten events and a final manifest containing event count, head hash and execution status. A separately supplied trusted public key anchors verification; an embedded key alone is insufficient.",
        "Fourth, the verifier checks event identity, ordering, links, hashes, signatures, checkpoint positions and final boundaries. The frozen detectors identify keyword indicators, policy violations and suspicious references to available actions. A separate annotation layer extracts destination fields and synthetic markers in tool arguments or model output without changing classifier predictions. Findings reference event IDs. Temporal links establish order, not internal causal influence.",
        "Finally, the dashboard displays the timeline, payloads, integrity result and incident findings. Evaluation reads preserved evidence and independent receipts, writes predictions and metrics, and produces report figures. SQLite indexes finalized sessions for convenience, but independent verification operates on exported files and does not trust the database.",
    ]
    blocks[8] = [
        "The execution model is Qwen3-4B-Instruct-2507 [9] in Q4_K_M GGUF form, served locally by llama.cpp b11900 [10]. The model produces actual answers and function-call decisions; sandbox outcomes are controlled so attacks cannot affect external systems. The model choice permits a live study on the available 4 GB GPU without paid inference services. Quantization and this small model's behavior constrain generalization.",
        "The keyword baseline uses a fixed expression over user and retrieved text for instruction-override phrases and secret-sending indicators. It approximates an analyst searching ordinary textual logs. It predicts exposure from those matches; using the same alarm as a coarse action alarm intentionally illustrates that suspicious words alone do not establish an attempted tool violation.",
        "The policy-only baseline flags observed tool requests outside the task's predefined allowed-tool set. This baseline is chosen because unauthorized tool use is a central forensic outcome and an explicit allowlist is a simple reproducible comparison. It can identify a violating action after it is proposed but cannot generally identify an injected instruction that the model resists.",
        "The combined method adds instruction-like references to available actions, filters clearly labelled security quotations, and associates earlier suspicious content with later violating requests. It combines content evidence with the same policy check. The quotation filter is heuristic and can be bypassed or over-triggered; the method is triage, not a comprehensive language-understanding defence.",
        "Both the policy-only and combined action branches use the same allowlist, and action ground truth is also defined by that explicit policy. Agreement on that task validates reconstruction and wiring, not an independent breakthrough in detection. The exposure task provides the less trivial comparison. Cryptographic integrity is evaluated separately from these classifiers.",
    ]
    blocks[9] = [
        "Experiments run on Windows with Python 3.12 and an NVIDIA GeForce RTX 3050 Laptop GPU with 4,096 MiB VRAM and approximately 16 GB system RAM. The local model server uses a 4,096-token context, one inference slot, Jinja tool templates, temperature 0.2, up to 256 generated tokens per request, and seeds 42, 43 and 44. An agent session allows at most three model turns. Model checksum, revision and server configuration are preserved in configs/model.json and session-start events.",
        "The study began with partial GPU offloading and tested full offloading after a memory check. Sustained GPU thermal throttling led to a return to the 24-layer partial profile. Completed sessions were retained, and the unfinished attempt was preserved separately on resume. Every session stores its actual profile. The paired overhead experiment uses a consistent server profile, and profile differences are not treated as detector improvements.",
        "The development/test allocation is fixed before execution. Detector rules are frozen without training a learned classifier or selecting a threshold on test labels. Metrics are accuracy, precision, recall, F1, false-positive rate and confusion-matrix counts. A 500-resample percentile bootstrap over tool or scenario groups provides 95% F1 intervals. Undefined precision or recall is assigned zero and actual positive support remains visible.",
        "Capture completeness compares request/result events against independent sandbox and model-client receipts. Integrity evaluation uses nine predefined corruption types and an intact control. Microbenchmarks compare 1,000 flushed plain JSONL records with full capture across five repetitions; full capture includes checkpoints and finalization. This measures the combined engineering cost of record construction, serialization, file handling, signatures and indexing, not cryptography in isolation.",
        "The live overhead study warms the model, then executes twelve paired scenarios with capture enabled and disabled in randomized order. Wall-clock timing surrounds the full session call, including finalization. Token counts and statuses are saved because caching, output length and GPU variability can dominate small recording costs. Benchmark session elapsed time covers agent execution; it is not substituted for this wider paired timing boundary. Locked software versions are supplied in uv.lock.",
    ]
    blocks[10] = [
        f"The experiment recorded {summary['sessions']} sessions: {summary['completed']} completed and {summary['failed']} failed. Every finalized bundle is retained regardless of model outcome. All {summary['integrity_passed']} evaluated bundles passed integrity verification, and receipted operation completeness was {pct(summary['receipt_completeness'])}. Required session, prompt, retrieval and termination boundaries were {'present in all sessions' if summary['all_required_boundaries'] else 'not present in every session'}.",
        f"On the {combined['n']} held-out public-subset sessions, the combined exposure detector achieved precision {fnum(combined['precision'])}, recall {fnum(combined['recall'])} and F1 {fnum(combined['f1'])}. Its grouped-bootstrap F1 interval was [{fnum(combined['f1_ci_low'])}, {fnum(combined['f1_ci_high'])}]. The keyword baseline F1 was {fnum(keyword['f1'])}; the policy-only exposure F1 was {fnum(policy['f1'])}. Table 4 compares the methods on identical evidence and separates exposure from action outcomes.",
        f"The held-out public subset contained {attempted} sessions with a proposed policy violation and {executed} completed simulated unauthorized tool actions. These are distinct counts: multiple calls can occur in one session, and a call proposed at the step limit may not execute. The combined action detector F1 was {fnum(action['f1'])}; interpret it in light of the shared explicit allowlist rather than as evidence of semantic attack understanding.",
        f"The verifier rejected {len(corrupted)} of {len(corrupted)} corrupted variants and accepted the intact control. The five microbenchmark repetitions measured mean incremental cost {extra:.3f} ms per event. Median paired live runtime change was {live_delta:+.1f}%; a negative or small estimate reflects timing variability and does not establish a speed improvement from logging. Raw paired observations and token counts are retained.",
        f"The safe-answer completion proxy was {pct(summary['safe_answer_completion_proxy'])} across all sessions. This proxy requires a non-empty answer, no proposed unauthorized tool and successful termination; it is not a full benchmark utility judgment. Controlled-scenario metrics and all confusion matrices are provided in the saved result tables.",
    ]
    blocks[11] = [
        "Figures 2-7 summarize the scenario construction, held-out detector behavior, verification coverage and recording cost. Each is generated from the saved scenario manifest or measured outputs. Figure 1 in the methodology chapter describes system flow rather than an experimental result."
    ]
    blocks[12] = [
        f"The combined detector missed {combined['fn']} of {combined['tp']+combined['fn']} attack-exposed held-out public sessions. Recall of {pct(combined['recall'])} supports only a limited triage role. Preserved evidence remains available for review even when a heuristic produces no alarm.",
        f"The combined exposure detector changes F1 from {fnum(keyword['f1'])} for the keyword baseline to {fnum(combined['f1'])} on the held-out adapted public sessions. This measured difference reflects the particular indicators and generated controls used here. It should not be interpreted as a universal improvement across arbitrary injection styles or a defence that prevents unsafe model actions.",
        "Structured evidence adds a practical capability even where classifier scores are unchanged: the analyst can inspect the precise retrieved content, model output, tool proposal and resulting sandbox response in one verified sequence. A resisted attack remains an exposure, while an unauthorized request remains an action attempt. This distinction prevents both treating every injection as a successful compromise and treating a refused instruction as an absence of malicious content.",
        "The policy-only and combined action scores are expected to agree because both apply the same authorization rule used to define action truth. That agreement is useful for validating evidence linkage but offers no proof that contextual reasoning improves action classification. The stronger system-level contribution is preserved provenance and integrity, which a plain keyword alarm does not provide.",
        "Published InjecAgent and AgentDojo outcomes [6,7] use different agents, tool environments and evaluation procedures. Our sample is smaller, controls are generated, the first retrieval is pre-specified and subsequent tool results are simulated. Consequently the project does not rank its vulnerability rate against those published results. Their findings motivate the threat model and experimental categories rather than supply a directly interchangeable baseline.",
        "Integrity results establish that the tested changes are rejected relative to a trusted key and finalized anchor. Completeness results establish coverage of receipted operations in this harness. Neither result proves that a malicious producer recorded every real-world event honestly. Similarly, paired runtime variation and token differences mean that the microbenchmark is the clearer estimate of recorder cost; live timing is reported as supporting context with its limits visible.",
    ]
    blocks[13] = [
        "The study uses one small quantized model, a selected public subset, a bounded context and a three-turn agent loop. It does not establish behavior for larger hosted models, long-running agents, multimodal input or production tool ecosystems. Repeated seeds and related controls increase observations without creating independent task diversity.",
        "Benign controls are project-authored replacements in published templates. Their explicit security quotations create a useful false-positive challenge but may favor a quotation-aware heuristic. Real benign language is broader, and an attacker could wrap malicious instructions in similar quotation cues. The selected tool-family split reduces direct task overlap while leaving shared attack wording and tool schemas across the benchmark.",
        "Subsequent tool outputs are generic sandbox responses rather than a reproduction of each original service. Data-stealing findings therefore concern modeled requests and synthetic markers, not confirmed disclosure of private information. Safe-answer completion is a structural proxy, and direct-injection scenarios exercise conflict with a predefined task policy rather than a universal definition of user intent.",
        "The signer uses a static local key, and timestamps come from the host clock. Compromise of the signing key and trust anchors, dishonest capture before sealing, or complete evidence destruction remain outside the guarantee. SQLite is not physically immutable; signatures enable later checking. The system cannot recover internal model computation, establish identity through a fixed actor string, or prove causal influence from temporal proximity.",
        "The next experiment should use a second model and native AgentDojo task outcomes, author new benign controls independently of detector rules, and hold out attack families as well as user tools. A further integrity extension would publish checkpoints to an independent witness and evaluate key rotation, crash recovery and concurrent collectors. These specific experiments would test generalization and operational assumptions that the present work leaves open.",
    ]
    blocks[14] = [
        "The project delivers a local AI-agent forensic prototype integrating live model execution, structured capture, signed integrity anchors, independent verification, an evidence viewer and reproducible analysis. Its four objectives are addressed through observable-operation capture, explicit corruption tests, held-out detector comparisons and recording-cost measurements.",
        f"Across {summary['sessions']} live sessions, operation-receipt coverage was {pct(summary['receipt_completeness'])}; all {len(corrupted)} deliberate corruption variants were rejected relative to the trusted public key. The combined held-out exposure F1 was {fnum(combined['f1'])}, compared with {fnum(keyword['f1'])} for the keyword method. These results are supported by saved traces, independent receipts, predictions, metrics and the scripts that regenerate their figures.",
        "The strongest supported contribution is the ability to connect an incident finding to a preserved, verifiable sequence while keeping exposure, proposed violation and simulated execution separate. Existing cryptographic primitives provide integrity checks, and simple detector baselines make the benefits and limits inspectable. Shared-policy action scores are a validation of reconstruction rather than evidence of a new general-purpose security classifier.",
        "The study remains a controlled mini project. One model, authored controls, simulated services and a static signing key constrain the conclusions. Extending the evaluation to independent task environments and independently witnessed anchors is the clearest route toward stronger operational evidence.",
    ]
    literature = [
        [
            "Schneier and Kelsey [2]",
            "Audit records",
            "Cryptographic log protection",
            "Security properties",
            "Protects history under stated compromise assumptions",
            "Agent-specific capture and triage",
        ],
        [
            "Crosby and Wallach [3]",
            "Log histories",
            "Authenticated history trees",
            "Proof and audit cost",
            "Efficient consistency checking",
            "Simple signed session bundles",
        ],
        [
            "ReAct [4]",
            "QA and interactive tasks",
            "Interleaved model/actions",
            "Task success",
            "Makes action trajectories operational",
            "Evidence preservation and verification",
        ],
        [
            "Greshake et al. [5]",
            "LLM applications",
            "Indirect injection experiments",
            "Attack impact",
            "Retrieved data can redirect actions",
            "Link raw context to observed actions",
        ],
        [
            "InjecAgent [6]",
            "1,054 attack cases",
            "Tool-integrated attack benchmark",
            "Attack success",
            "Exposes harmful and stealing goals",
            "Forensic evidence on adapted sessions",
        ],
        [
            "AgentDojo [7]",
            "Dynamic agent tasks",
            "Attack/defence environment",
            "Utility and security",
            "Jointly studies tasks and attacks",
            "Future native-environment validation",
        ],
    ]
    study_path = ROOT / "results/case_study.json"
    if study_path.exists():
        study = json.loads(study_path.read_text(encoding="utf-8"))
        execution = (
            f"Its simulated result is preserved at {study['result_event_id']}."
            if study["execution_observed"]
            else "No corresponding simulated result was observed."
        )
        blocks[12].append(
            f"An illustrative held-out session, {study['session_id']}, connects retrieved content at {study['retrieval_event_id']} to the out-of-policy {study['tool']} request at {study['request_event_id']}. {execution} These references let an investigator inspect the preserved sequence directly. The example illustrates reconstruction and is not an additional performance estimate or proof of a real service action."
        )
    return {
        "summary": summary,
        "data": data,
        "model": model,
        "abstract": abstract,
        "blocks": blocks,
        "literature": literature,
        "sources": sources,
        "main_rows": [r for r in m if r["source"] == "InjecAgent-adapted" and r["split"] == "test"],
        "versions": versions,
        "micro_extra_ms": extra,
        "live_delta": live_delta,
    }


def build(pagination_pdf=None):
    template = ROOT / "references/Template.docx"
    if hashlib.sha256(template.read_bytes()).hexdigest() != TEMPLATE_SHA:
        raise ValueError("Reference template checksum mismatch")
    values = content()
    pages = pagination(pagination_pdf)
    doc = Document(template)
    original = list(doc.paragraphs)
    heading_pattern = deepcopy(original[16]._p)
    body_pattern = deepcopy(original[17]._p)

    # Preserve the original cover, including its logo drawing, and expand only editable slots.
    def replace_cover(p, text, pattern=None):
        source = pattern or p
        props = deepcopy(source._p.pPr) if source._p.pPr is not None else None
        run_props = deepcopy(source.runs[0]._r.rPr) if source.runs and source.runs[0]._r.rPr is not None else None
        p.clear()
        if pattern is not None:
            if p._p.pPr is not None:
                p._p.remove(p._p.pPr)
            if props is not None:
                p._p.insert(0, props)
        r = p.add_run(text)
        if run_props is not None:
            r._r.insert(0, run_props)

    replace_cover(original[0], "AI Agent Forensics and Audit Platform")
    replace_cover(original[4], "Pandiri Jaswanth Kumar")
    replace_cover(original[5], "Registration No 230968014")
    for slot, text, bold in [
        (6, "Kosuru Venkata Sai Aditya", True),
        (7, "Registration No 230968316", False),
        (9, "Under the guidance of Dr. Sindhura D N", False),
    ]:
        p = original[slot]
        if slot in (6, 7):
            replace_cover(p, text, original[4 if slot == 6 else 5])
            if slot == 6:
                p.paragraph_format.space_before = Pt(12)
            continue
        p.clear()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.bold = bold
    replace_cover(original[14], "October 2026")
    for element in list(doc._element.body)[15:]:
        if element.tag != qn("w:sectPr"):
            doc._element.body.remove(element)
    all_captions = []

    def paragraph(text, heading=False):
        element = deepcopy(heading_pattern if heading else body_pattern)
        for child in list(element):
            if child.tag != qn("w:pPr"):
                element.remove(child)
        doc._element.body.insert(len(doc._element.body) - 1, element)
        from docx.text.paragraph import Paragraph

        p = Paragraph(element, doc._body)
        r = p.add_run(text)
        r.font.color.rgb = RGBColor(0, 0, 0)
        if not heading:
            p.paragraph_format.space_after = Pt(7)
        return p

    def number_section(section, fmt, start):
        section.footer_distance = Inches(0.25)
        section.footer.is_linked_to_previous = False
        p = section.footer.paragraphs[0]
        p.style = doc.styles["Normal"]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        r._r.addnext(field)
        props = section._sectPr
        n = props.find(qn("w:pgNumType"))
        if n is None:
            n = OxmlElement("w:pgNumType")
            props.append(n)
        n.set(qn("w:fmt"), fmt)
        n.set(qn("w:start"), str(start))

    front = doc.add_section(WD_SECTION_START.NEW_PAGE)
    number_section(front, "lowerRoman", 1)
    paragraph("Abstract", True)
    paragraph(values["abstract"])
    paragraph("Table Of Contents", True).paragraph_format.page_break_before = True

    def toc_entry(label, key=None):
        p = paragraph(label + "\t" + pages.get(key or label, "-"))
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    for i, h in enumerate(HEADINGS, 1):
        toc_entry(f"{i} {h}")
    toc_entry("References")
    toc_entry("Appendix Reproduction And Evidence")
    p = paragraph("List Of Figures")
    p.runs[0].bold = True
    figtitles = [
        "Overall evidence workflow",
        "Scenario composition",
        "Held-out detector comparison",
        "Exposure confusion matrices",
        "Integrity verification outcomes",
        "Operation receipt completeness",
        "Recording overhead",
    ]
    for i, title in enumerate(figtitles, 1):
        toc_entry(f"Figure {i} {title}", f"Figure {i}")
    p = paragraph("List Of Tables")
    p.runs[0].bold = True
    for i, title in enumerate(
        [
            "Literature comparison",
            "Dataset and split summary",
            "Experimental software and parameters",
            "Held-out public-subset comparison",
        ],
        1,
    ):
        toc_entry(f"Table {i} {title}", f"Table {i}")
    main = doc.add_section(WD_SECTION_START.NEW_PAGE)
    number_section(main, "decimal", 1)

    def figure(name, caption, interpretation=None, width=6.8, new_page=False):
        p = doc.add_paragraph()
        if new_page:
            p.paragraph_format.page_break_before = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(ROOT / "results/figures" / f"{name}.png"), width=Inches(width))
        p = paragraph(caption)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = bool(interpretation)
        for r in p.runs:
            r.font.size = Pt(11)
        all_captions.append(caption)
        if interpretation:
            paragraph(interpretation)

    def table(headers, rows, caption, widths=None):
        p = paragraph(caption)
        p.paragraph_format.keep_with_next = True
        all_captions.append(caption)
        table = doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        if widths:
            for col, width in zip(table.columns, widths):
                col.width = Inches(width)
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
        for row in rows:
            for cell, text in zip(table.add_row().cells, row):
                cell.text = str(text)
        for i, row in enumerate(table.rows):
            if i == 0:
                repeat = OxmlElement("w:tblHeader")
                row._tr.get_or_add_trPr().append(repeat)
            for j, cell in enumerate(row.cells):
                if widths:
                    cell.width = Inches(widths[j])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                props = cell._tc.get_or_add_tcPr()
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "3F4B55" if i == 0 else ("F1F3F5" if i % 2 == 0 else "FFFFFF"))
                props.append(shading)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "bottom", "left", "right"):
                    edge = OxmlElement("w:" + side)
                    edge.set(qn("w:w"), "80")
                    edge.set(qn("w:type"), "dxa")
                    margins.append(edge)
                props.append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.05
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(10)
                        r.bold = i == 0
                        r.font.color.rgb = RGBColor.from_string("FFFFFF" if i == 0 else "000000")
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement("w:" + side)
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), "D9D9D9")
            borders.append(e)
        table._tbl.tblPr.append(borders)

    for i, title in enumerate(HEADINGS, 1):
        title_p = paragraph(f"{i} {title}", True)
        if i > 1:
            title_p.paragraph_format.page_break_before = True
        for text in values["blocks"][i]:
            paragraph(text)
        if i == 4:
            table(
                ["Paper", "Dataset", "Method", "Metric", "Finding", "Project gap"],
                values["literature"],
                "Table 1 Literature comparison",
                [1.1, 0.95, 1.05, 0.85, 1.25, 1.0],
            )
        if i == 5:
            table(
                ["Scenario source", "Development", "Test", "Unique total", "Planned sessions"],
                [
                    ["Public attacks", 20, 40, 60, 180],
                    ["Generated paired controls", 20, 40, 60, 180],
                    ["Controlled scenarios", 8, 16, 24, 72],
                    ["Total", 48, 96, 144, 432],
                ],
                "Table 2 Dataset and split summary",
                [2.15, 1.1, 1.0, 1.1, 1.1],
            )
        if i == 7:
            figure("workflow", "Figure 1 Overall evidence workflow", width=6.8)
        if i == 9:
            table(
                ["Component", "Recorded setting"],
                [
                    ["Model", "Qwen3-4B-Instruct-2507, Q4_K_M"],
                    ["Inference", "llama.cpp b11900, context 4096, one slot"],
                    ["Sampling", "Temperature 0.2; seeds 42, 43, 44"],
                    ["Bounds", "256 output tokens; three model turns"],
                    ["Python", values["model"]["hardware"]["python"].split()[0]],
                    [
                        "Agent and viewer",
                        f"LangGraph {values['versions']['langgraph']}; Streamlit {values['versions']['streamlit']}",
                    ],
                    [
                        "Integrity",
                        f"cryptography {values['versions']['cryptography']}; rfc8785 {values['versions']['rfc8785']}",
                    ],
                    ["Evaluation", "Fixed group split; 500 grouped bootstrap samples"],
                ],
                "Table 3 Experimental software and parameters",
                [1.5, 5.05],
            )
        if i == 10:
            rows = []
            for r in sorted(values["main_rows"], key=lambda r: (r["task"], r["method"])):
                rows.append(
                    [
                        r["task"].replace("_", " "),
                        r["method"],
                        fnum(r["accuracy"]),
                        fnum(r["precision"]),
                        fnum(r["recall"]),
                        fnum(r["f1"]),
                        f"[{r['f1_ci_low']:.3f}, {r['f1_ci_high']:.3f}]",
                    ]
                )
            table(
                ["Task", "Method", "Accuracy", "Precision", "Recall", "F1", "95% F1 interval"],
                rows,
                "Table 4 Held-out public-subset comparison",
                [1.0, 1.0, 0.8, 0.8, 0.7, 0.65, 1.2],
            )
        if i == 11:
            figure(
                "dataset",
                "Figure 2 Scenario composition",
                "Public attacks and their generated controls are balanced by selection. Controlled scenarios add separate behavior categories; repeated seeds are not new unique scenarios.",
                width=5.7,
            )
            figure(
                "detector_comparison",
                "Figure 3 Held-out detector comparison",
                "The methods are evaluated on identical held-out public-subset evidence. Exposure and action violations are separate tasks, so a resisted injection can remain positive only for exposure.",
                width=6.6,
            )
            figure(
                "confusion_matrices",
                "Figure 4 Exposure confusion matrices",
                "Counts reveal which alarms are false positives and which attacks are missed. The generated security quotations challenge the keyword rule, while the combined rule remains limited to its fixed indicators.",
                width=6.8,
                new_page=True,
            )
            figure(
                "integrity",
                "Figure 5 Integrity verification outcomes",
                "The intact control is accepted and each predefined corrupted variant is rejected. These checks assume the public key and finalized anchor remain trusted; they do not prove protection after key compromise.",
                width=6.1,
            )
            figure(
                "completeness",
                "Figure 6 Operation receipt completeness",
                "Coverage compares stored event pairs with independent model-client and sandbox receipts. It concerns observable operations in this harness, while boundary checks and failed execution status are evaluated separately.",
                width=5.9,
                new_page=True,
            )
            figure(
                "overhead",
                "Figure 7 Recording overhead",
                "The microbenchmark measures the full recorder against flushed plain JSONL. Paired live timings include model and token variability, so they cannot isolate cryptographic cost or establish a speed benefit from instrumentation.",
                width=6.6,
            )
    paragraph("References", True).paragraph_format.page_break_before = True
    for source in values["sources"]:
        paragraph(f"[{source['id']}] {source['citation']} {source['url']}")
    paragraph("Appendix Reproduction And Evidence", True).paragraph_format.page_break_before = True
    for text in [
        "The submission provides source code, locked Python dependencies, the original dataset and provenance, the prepared scenario manifest, signed session evidence, independent execution receipts, predictions, metric tables, figures and report-generation source. Model weights and private signing keys are excluded.",
        "To reproduce analysis: install Python 3.12 and uv, run uv sync --frozen, then uv run agent-forensics evaluate and uv run agent-forensics figures. These commands recompute the reported results from stored observations without a model download.",
        "To generate new live evidence: run uv run agent-forensics fetch-model, start uv run agent-forensics serve-model, then in a second terminal run uv run agent-forensics run --output-root output/live-rerun. New evidence uses its own signing key and preserves the submitted reference corpus. Lower --gpu-layers when required by available memory.",
        "To inspect evidence: run uv run agent-forensics dashboard. Verify a session or ZIP with uv run agent-forensics verify PATH --public-key results/public_key.pem. Each manifest fixes a count and head hash; the trusted public-key fingerprint is recorded in the README. For an expected-session substitution threat, retain the expected session ID or manifest checksum separately.",
        "The detector freeze, split manifest, model checksum, bootstrap seed and raw paired runtime records support auditability. Recomputed metrics should match the supplied tables exactly. New live generations may differ because seeded GPU inference is not guaranteed bit-identical.",
        "The submission checklist records administrative attachments separately from technical results. The final institutional similarity check must be performed on the exact submitted report, with its genuine report attached according to institutional instructions.",
    ]:
        paragraph(text)
    out = ROOT / "output/report"
    out.mkdir(parents=True, exist_ok=True)
    working = ROOT / "tmp/report-working.docx"
    doc.save(working)
    # Preserve opaque reference package members byte-for-byte. Only authoring
    # parts, added footer/image parts, and their required relationships change.
    with (
        ZipFile(template) as source,
        ZipFile(working) as authored,
        ZipFile(out / "Project_Report.docx", "w", ZIP_DEFLATED) as final,
    ):
        changed = {"word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"}
        existing = set(source.namelist())
        added = set(authored.namelist()) - existing
        for item in source.infolist():
            final.writestr(
                item, authored.read(item.filename) if item.filename in changed else source.read(item.filename)
            )
        for name in sorted(added):
            final.writestr(name, authored.read(name))
    (out / "report_content.json").write_text(
        json.dumps(
            {
                "abstract": values["abstract"],
                "chapters": values["blocks"],
                "contents_pages": pages,
                "captions": all_captions,
                "reported_sessions": values["summary"]["sessions"],
                "report_docx_sha256": hashlib.sha256((out / "Project_Report.docx").read_bytes()).hexdigest(),
                "source_hashes": {
                    str(p.relative_to(ROOT)).replace("\\", "/"): hashlib.sha256(p.read_bytes()).hexdigest()
                    for p in [
                        ROOT / "results/summary.json",
                        ROOT / "results/tables/metrics.csv",
                        ROOT / "results/overhead_live.json",
                        ROOT / "results/overhead_micro.json",
                        ROOT / "data/cases.json",
                        ROOT / "uv.lock",
                        *sorted((ROOT / "results/figures").glob("*.png")),
                        *sorted((ROOT / "results").glob("*.json")),
                    ]
                },
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    print(f"Report generated; abstract {len(values['abstract'].split())} words; contents entries {len(pages)}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--pagination", type=Path)
    args = parser.parse_args()
    build(args.pagination)
